"""
Tests para el módulo exportador_resultados.py
"""

import json
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import load_workbook


def test_exportador_import_does_not_import_polars():
    sys.modules.pop("s50exportador_resultados", None)
    sys.modules.pop("polars", None)

    __import__("s50exportador_resultados")

    assert "polars" not in sys.modules


class TestExportadorResultadosInit:
    """Tests para el constructor de ExportadorResultados"""

    def test_init_creates_output_directory(self, tmp_path):
        """Test que el constructor crea el directorio de salida"""
        from s50exportador_resultados import ExportadorResultados

        output_dir = tmp_path / "test_output"
        exportador = ExportadorResultados(str(output_dir))

        assert exportador.directorio_salida == output_dir
        assert output_dir.exists()
        assert output_dir.is_dir()


class TestExportar:
    """Tests para el método exportar"""

    @pytest.fixture
    def sample_data(self):
        """Fixture con datos de ejemplo"""
        return [
            {"id": 1, "nombre": "Juan", "edad": 30},
            {"id": 2, "nombre": "María", "edad": 25},
        ]

    @pytest.fixture
    def exportador(self, tmp_path):
        """Fixture para exportador con directorio temporal"""
        from s50exportador_resultados import ExportadorResultados

        return ExportadorResultados(str(tmp_path))

    def test_exportar_with_format_txt(self, exportador, sample_data):
        """Test exportación a formato TXT"""
        resultado = exportador.exportar(sample_data, formato="txt", nombre_archivo="test")

        assert Path(resultado).exists()
        assert Path(resultado).suffix == ".txt"

    def test_exportar_with_format_csv(self, exportador, sample_data):
        """Test exportación a formato CSV"""
        resultado = exportador.exportar(sample_data, formato="csv", nombre_archivo="test")

        assert Path(resultado).exists()
        assert Path(resultado).suffix == ".csv"

        # Verificar contenido
        with open(resultado, "r", encoding="utf-8-sig") as f:
            contenido = f.read()
            assert "id" in contenido
            assert "Juan" in contenido
            assert "María" in contenido

    def test_exportar_with_format_json(self, exportador, sample_data):
        """Test exportación a formato JSON"""
        resultado = exportador.exportar(sample_data, formato="json", nombre_archivo="test")

        assert Path(resultado).exists()
        assert Path(resultado).suffix == ".json"

        # Verificar contenido
        with open(resultado, "r", encoding="utf-8") as f:
            data = json.load(f)
            assert "metadatos" in data
            assert "datos" in data
            assert data["metadatos"]["total_registros"] == 2
            assert len(data["datos"]) == 2

    def test_exportar_with_format_xml(self, exportador, sample_data):
        """Test exportación a formato XML"""
        resultado = exportador.exportar(sample_data, formato="xml", nombre_archivo="test")

        assert Path(resultado).exists()
        assert Path(resultado).suffix == ".xml"

        # Verificar contenido
        with open(resultado, "r", encoding="utf-8") as f:
            contenido = f.read()
            assert "<resultados>" in contenido
            assert "<metadatos>" in contenido
            assert "<datos>" in contenido
            assert "<registro" in contenido

    def test_exportar_with_format_excel(self, exportador, sample_data):
        """Test exportación a formato Excel"""
        resultado = exportador.exportar(
            sample_data, formato="excel", nombre_archivo="test", abrir_archivo=False
        )

        assert Path(resultado).exists()
        assert Path(resultado).suffix == ".xlsx"

        wb = load_workbook(resultado)
        ws_datos = wb["Datos"]
        assert [cell.value for cell in ws_datos[1]] == ["id", "nombre", "edad"]
        assert [cell.value for cell in ws_datos[2]] == [1, "Juan", 30]

        ws_metadatos = wb["Metadatos"]
        assert ws_metadatos[3][1].value == 2

    def test_exportar_with_format_excel_mixed_values(self, exportador):
        """Excel se exporta sin Polars aunque haya valores mixtos."""
        sample_data = [
            {"id": 1, "fecha": "texto"},
            {"id": 2, "fecha": datetime(2001, 1, 1, 0, 0, 0)},
        ]

        resultado = exportador.exportar(
            sample_data, formato="excel", nombre_archivo="test", abrir_archivo=False
        )

        wb = load_workbook(resultado)
        ws_datos = wb["Datos"]
        assert ws_datos[2][1].value == "texto"
        assert ws_datos[3][1].value == datetime(2001, 1, 1, 0, 0, 0)

    def test_exportar_with_unsupported_format(self, exportador, sample_data):
        """Test con formato no soportado"""
        with pytest.raises(ValueError, match="Formato no soportado"):
            exportador.exportar(sample_data, formato="invalid")

    def test_exportar_with_empty_data(self, exportador):
        """Test con datos vacíos"""
        with pytest.raises(ValueError, match="No hay datos para exportar"):
            exportador.exportar([], formato="txt")

    def test_exportar_with_compression(self, exportador, sample_data):
        """Test exportación con compresión"""
        resultado = exportador.exportar(
            sample_data, formato="txt", nombre_archivo="test", comprimir=True, abrir_archivo=False
        )

        # Verificar que se creó un ZIP
        assert Path(resultado).exists()
        assert Path(resultado).suffix == ".zip"

        # Verificar que el archivo original no existe
        original_file = Path(resultado).with_suffix(".txt")
        assert not original_file.exists()

        # Verificar contenido del ZIP
        with zipfile.ZipFile(resultado, "r") as zf:
            assert len(zf.namelist()) == 1
            assert zf.namelist()[0].endswith(".txt")

    @patch("s50exportador_resultados.subprocess.Popen")
    def test_exportar_with_template(self, mock_popen, exportador, sample_data, tmp_path):
        """Test exportación con plantilla personalizada"""
        # Crear plantilla
        template_file = tmp_path / "template.txt"
        template_content = """
        Registro: {NUMERO}
        Fecha: {FECHA}
        Nombre: {NOMBRE}
        Edad: {EDAD}
        """
        template_file.write_text(template_content, encoding="utf-8")

        # Datos adaptados a la plantilla
        data = [{"NOMBRE": "Juan", "EDAD": 30}, {"NOMBRE": "María", "EDAD": 25}]

        resultado = exportador.exportar(
            data, formato="txt", plantilla=str(template_file), abrir_archivo=False
        )

        # Verificar que el archivo se creó
        assert Path(resultado).exists()

        # Verificar contenido
        with open(resultado, "r", encoding="utf-8") as f:
            contenido = f.read()
            assert "Registro: 1" in contenido
            assert "Registro: 2" in contenido
            assert "Nombre: Juan" in contenido
            assert "Nombre: María" in contenido
            assert "Edad: 30" in contenido

    @patch("builtins.print")
    @patch("s50exportador_resultados.subprocess.Popen")
    def test_exportar_with_missing_template_creates_base_template(
        self, mock_popen, mock_print, exportador, tmp_path
    ):
        data = [{"NOMBRE": "Juan", "EDAD": 30}]
        template_file = tmp_path / "faltante.txt"

        resultado = exportador.exportar(
            data, formato="txt", plantilla=str(template_file), abrir_archivo=False
        )

        assert Path(resultado).exists()
        assert template_file.exists()
        contenido_plantilla = template_file.read_text(encoding="utf-8")
        assert "NOMBRE\tEDAD" in contenido_plantilla
        assert "{NOMBRE}\t{EDAD}" in contenido_plantilla
        print_calls = [str(call) for call in mock_print.call_args_list]
        assert any("Plantilla no encontrada" in call for call in print_calls)


class TestExportarTxt:
    """Tests específicos para exportación TXT"""

    @pytest.fixture
    def exportador(self, tmp_path):
        from s50exportador_resultados import ExportadorResultados

        return ExportadorResultados(str(tmp_path))

    @pytest.fixture
    def sample_data(self):
        return [
            {"id": 1, "nombre": "Juan Pérez"},
            {"id": 2, "nombre": "María López"},
        ]

    def test_exportar_txt_creates_file(self, exportador, sample_data):
        """Test que crea archivo TXT"""
        resultado = exportador._exportar_txt(sample_data, "test")

        assert resultado.exists()
        assert resultado.suffix == ".txt"

    def test_exportar_txt_content(self, exportador, sample_data):
        """Test contenido del archivo TXT"""
        resultado = exportador._exportar_txt(sample_data, "test")

        with open(resultado, "r", encoding="utf-8") as f:
            contenido = f.read()
            assert "id" in contenido
            assert "nombre" in contenido
            assert "Juan Pérez" in contenido
            assert "María López" in contenido


class TestExportarCsv:
    """Tests específicos para exportación CSV"""

    @pytest.fixture
    def exportador(self, tmp_path):
        from s50exportador_resultados import ExportadorResultados

        return ExportadorResultados(str(tmp_path))

    @pytest.fixture
    def sample_data(self):
        return [
            {"id": 1, "nombre": "Juan", "edad": 30},
            {"id": 2, "nombre": "María", "edad": 25},
        ]

    def test_exportar_csv_creates_file(self, exportador, sample_data):
        """Test que crea archivo CSV"""
        resultado = exportador._exportar_csv(sample_data, "test")

        assert resultado.exists()
        assert resultado.suffix == ".csv"

    def test_exportar_csv_content(self, exportador, sample_data):
        """Test contenido del archivo CSV"""
        resultado = exportador._exportar_csv(sample_data, "test")

        with open(resultado, "r", encoding="utf-8-sig") as f:
            contenido = f.read()
            # Verificar cabecera
            assert "id,nombre,edad" in contenido
            # Verificar datos
            assert "1,Juan,30" in contenido or '"1","Juan","30"' in contenido


class TestExportarJson:
    """Tests específicos para exportación JSON"""

    @pytest.fixture
    def exportador(self, tmp_path):
        from s50exportador_resultados import ExportadorResultados

        return ExportadorResultados(str(tmp_path))

    @pytest.fixture
    def sample_data(self):
        return [
            {"id": 1, "nombre": "Juan"},
            {"id": 2, "nombre": "María"},
        ]

    def test_exportar_json_creates_file(self, exportador, sample_data):
        """Test que crea archivo JSON"""
        resultado = exportador._exportar_json(sample_data, "test")

        assert resultado.exists()
        assert resultado.suffix == ".json"

    def test_exportar_json_structure(self, exportador, sample_data):
        """Test estructura del archivo JSON"""
        resultado = exportador._exportar_json(sample_data, "test")

        with open(resultado, "r", encoding="utf-8") as f:
            data = json.load(f)

        assert "metadatos" in data
        assert "datos" in data
        assert "fecha_generacion" in data["metadatos"]
        assert "total_registros" in data["metadatos"]
        assert data["metadatos"]["total_registros"] == 2
        assert len(data["datos"]) == 2


class TestExportarXml:
    """Tests específicos para exportación XML"""

    @pytest.fixture
    def exportador(self, tmp_path):
        from s50exportador_resultados import ExportadorResultados

        return ExportadorResultados(str(tmp_path))

    @pytest.fixture
    def sample_data(self):
        return [
            {"id": 1, "nombre": "Juan"},
            {"id": 2, "nombre": "María"},
        ]

    def test_exportar_xml_creates_file(self, exportador, sample_data):
        """Test que crea archivo XML"""
        resultado = exportador._exportar_xml(sample_data, "test")

        assert resultado.exists()
        assert resultado.suffix == ".xml"

    def test_exportar_xml_structure(self, exportador, sample_data):
        """Test estructura del archivo XML"""
        resultado = exportador._exportar_xml(sample_data, "test")

        with open(resultado, "r", encoding="utf-8") as f:
            contenido = f.read()

        assert "<?xml" in contenido
        assert "<resultados>" in contenido
        assert "<metadatos>" in contenido
        assert "<datos>" in contenido
        assert "<registro" in contenido


class TestAplicarPlantillaTxt:
    """Tests para _aplicar_plantilla_txt"""

    @pytest.fixture
    def exportador(self, tmp_path):
        from s50exportador_resultados import ExportadorResultados

        return ExportadorResultados(str(tmp_path))

    @pytest.fixture
    def template_file(self, tmp_path):
        template_content = """
        Registro N°: {NUMERO}
        Fecha: {FECHA}
        Nombre: {NOMBRE}
        Edad: {EDAD}
        """
        template = tmp_path / "template.txt"
        template.write_text(template_content, encoding="utf-8")
        return template

    @pytest.fixture
    def sample_data(self):
        return [{"NOMBRE": "Juan", "EDAD": 30}, {"NOMBRE": "María", "EDAD": 25}]

    def test_aplicar_plantilla_txt_replace_numero(self, exportador, template_file, sample_data):
        """Test reemplazo de marcador {NUMERO}"""
        output_file = exportador.directorio_salida / "test.txt"
        resultado = exportador._aplicar_plantilla_txt(sample_data, output_file, str(template_file))

        with open(resultado, "r", encoding="utf-8") as f:
            contenido = f.read()

        assert "Registro N°: 1" in contenido
        assert "Registro N°: 2" in contenido

    def test_aplicar_plantilla_txt_replace_fecha(self, exportador, template_file, sample_data):
        """Test reemplazo de marcador {FECHA}"""
        output_file = exportador.directorio_salida / "test.txt"
        resultado = exportador._aplicar_plantilla_txt(sample_data, output_file, str(template_file))

        with open(resultado, "r", encoding="utf-8") as f:
            contenido = f.read()

        # Verificar formato de fecha
        assert "Fecha:" in contenido

    def test_aplicar_plantilla_txt_replace_fields(self, exportador, template_file, sample_data):
        """Test reemplazo de campos personalizados"""
        output_file = exportador.directorio_salida / "test.txt"
        resultado = exportador._aplicar_plantilla_txt(sample_data, output_file, str(template_file))

        with open(resultado, "r", encoding="utf-8") as f:
            contenido = f.read()

        assert "Nombre: Juan" in contenido
        assert "Nombre: María" in contenido
        assert "Edad: 30" in contenido
        assert "Edad: 25" in contenido


class TestComprimirArchivo:
    """Tests para _comprimir_archivo"""

    @pytest.fixture
    def exportador(self, tmp_path):
        from s50exportador_resultados import ExportadorResultados

        return ExportadorResultados(str(tmp_path))

    @pytest.fixture
    def test_file(self, tmp_path):
        file = tmp_path / "test.txt"
        file.write_text("Contenido de prueba", encoding="utf-8")
        return file

    def test_comprimir_archivo_creates_zip(self, exportador, test_file):
        """Test que crea archivo ZIP"""
        resultado = exportador._comprimir_archivo(test_file)

        assert resultado.exists()
        assert resultado.suffix == ".zip"

    def test_comprimir_archivo_removes_original(self, exportador, test_file):
        """Test que elimina archivo original"""
        exportador._comprimir_archivo(test_file)

        assert not test_file.exists()

    def test_comprimir_archivo_zip_contents(self, exportador, test_file):
        """Test contenido del ZIP"""
        resultado = exportador._comprimir_archivo(test_file)

        with zipfile.ZipFile(resultado, "r") as zf:
            assert len(zf.namelist()) == 1
            with zf.open(zf.namelist()[0]) as f:
                contenido = f.read().decode("utf-8")
                assert "Contenido de prueba" in contenido


class TestAbrirArchivo:
    """Tests para _abrir_archivo"""

    @pytest.fixture
    def exportador(self, tmp_path):
        from s50exportador_resultados import ExportadorResultados

        return ExportadorResultados(str(tmp_path))

    @patch("s50exportador_resultados.subprocess.Popen")
    def test_abrir_archivo_csv(self, mock_popen, exportador, tmp_path):
        """Test abrir archivo CSV"""
        csv_file = tmp_path / "test.csv"
        csv_file.write_text("test", encoding="utf-8")

        exportador._abrir_archivo(csv_file)

        mock_popen.assert_called_once()
        call_args = mock_popen.call_args[0][0]
        assert "notepad.exe" in call_args

    @patch("s50exportador_resultados.subprocess.Popen")
    def test_abrir_archivo_json(self, mock_popen, exportador, tmp_path):
        """Test abrir archivo JSON"""
        json_file = tmp_path / "test.json"
        json_file.write_text("{}", encoding="utf-8")

        exportador._abrir_archivo(json_file)

        mock_popen.assert_called_once()

    @patch("s50exportador_resultados.subprocess.Popen")
    def test_abrir_archivo_xml(self, mock_popen, exportador, tmp_path):
        """Test abrir archivo XML"""
        xml_file = tmp_path / "test.xml"
        xml_file.write_text("<root/>", encoding="utf-8")

        exportador._abrir_archivo(xml_file)

        mock_popen.assert_called_once()


class TestCrearPlantillaEjemplo:
    """Tests para crear_plantilla_ejemplo"""

    @pytest.fixture
    def exportador(self, tmp_path):
        from s50exportador_resultados import ExportadorResultados

        return ExportadorResultados(str(tmp_path))

    def test_crear_plantilla_ejemplo_creates_file(self, exportador):
        """Test que crea plantilla de ejemplo"""
        resultado = exportador.crear_plantilla_ejemplo("mi_plantilla.txt")

        assert resultado.exists()
        assert resultado.name == "mi_plantilla.txt"

    def test_crear_plantilla_ejemplo_content(self, exportador):
        """Test contenido de la plantilla de ejemplo"""
        resultado = exportador.crear_plantilla_ejemplo()

        with open(resultado, "r", encoding="utf-8") as f:
            contenido = f.read()

        assert "{NUMERO}" in contenido
        assert "{FECHA}" in contenido
        assert "{CAMPO1}" in contenido
        assert "{VALOR1}" in contenido

    def test_crear_plantilla_desde_formato_actual_content(self, exportador):
        resultado = exportador._crear_plantilla_desde_formato_actual(
            [{"CODIGO": "A1", "NOMBRE": "Cliente"}]
        )

        contenido = resultado.read_text(encoding="utf-8")
        assert "CODIGO\tNOMBRE" in contenido
        assert "{CODIGO}\t{NOMBRE}" in contenido


class TestLimpiarResultadosAntiguos:
    """Tests para limpiar_resultados_antiguos"""

    @pytest.fixture
    def exportador(self, tmp_path):
        from s50exportador_resultados import ExportadorResultados

        return ExportadorResultados(str(tmp_path))

    @pytest.fixture
    def old_file(self, exportador, tmp_path):
        """Archivo antiguo (más de 7 días)"""
        old_file = exportador.directorio_salida / "old.txt"
        old_file.write_text("old", encoding="utf-8")
        # Simular archivo antiguo
        old_time = datetime.now().timestamp() - (8 * 24 * 3600)  # 8 días atrás
        import os

        os.utime(old_file, (old_time, old_time))
        return old_file

    @pytest.fixture
    def new_file(self, exportador, tmp_path):
        """Archivo reciente"""
        new_file = exportador.directorio_salida / "new.txt"
        new_file.write_text("new", encoding="utf-8")
        return new_file

    def test_limpiar_resultados_antiguos_deletes_old_files(self, exportador, old_file, new_file):
        """Test que elimina archivos antiguos"""
        exportador.limpiar_resultados_antiguos(dias=7)

        assert not old_file.exists()
        assert new_file.exists()


class TestExportarResultadosFunction:
    """Tests para la función de conveniencia exportar_resultados"""

    @patch("s50exportador_resultados.ExportadorResultados")
    def test_exportar_resultados_function(self, mock_exportador_class, tmp_path):
        """Test función de conveniencia"""
        from s50exportador_resultados import exportar_resultados

        mock_exportador = MagicMock()
        mock_exportador.exportar.return_value = "resultados/test.txt"
        mock_exportador_class.return_value = mock_exportador

        datos = [{"id": 1, "nombre": "Juan"}]
        resultado = exportar_resultados(datos, formato="txt")

        mock_exportador_class.assert_called_once()
        mock_exportador.exportar.assert_called_once_with(datos, "txt", None)
